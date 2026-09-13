from datetime import datetime


class ExcelClassSectionImportService:

    REQUIRED_HEADERS = [
        "Tên môn",
        "Nhóm",
        "Tên giảng viên",
        "Học kỳ",
        "Năm học",
        "Ngày bắt đầu",
        "Ngày kết thúc",
    ]

    def read(
        self,
        file,
    ):

        try:

            import openpyxl

            workbook = openpyxl.load_workbook(
                file,
                data_only=True,
            )

            worksheet = workbook.active

        except Exception:

            raise ValueError("Không thể đọc file Excel.")

        headers = [cell.value for cell in worksheet[1]]

        normalized_headers = [
            str(header).strip() if header is not None else "" for header in headers
        ]

        if normalized_headers != self.REQUIRED_HEADERS:

            raise ValueError("File Excel không đúng định dạng cột.")

        rows = []

        for values in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if all(value is None for value in values):
                continue

            row = dict(
                zip(
                    self.REQUIRED_HEADERS,
                    values,
                )
            )

            row["Ngày bắt đầu"] = self._parse_date(row["Ngày bắt đầu"])

            row["Ngày kết thúc"] = self._parse_date(row["Ngày kết thúc"])
            print(
                f"[EXCEL] Dòng {worksheet.max_row}: "
                f"{row['Ngày bắt đầu']} -> "
                f"{row['Ngày kết thúc']}"
            )
            rows.append(row)

        if not rows:

            raise ValueError("File Excel không có dữ liệu.")

        return rows

    def _parse_date(
        self,
        value,
    ):

        if not value:

            return None

        if hasattr(
            value,
            "date",
        ):

            return value.date()

        if isinstance(
            value,
            str,
        ):

            value = value.strip()

            for date_format in [
                "%d/%m/%Y",
                "%Y-%m-%d",
            ]:

                try:

                    return datetime.strptime(
                        value,
                        date_format,
                    ).date()

                except ValueError:
                    continue

        raise ValueError(f"Ngày không hợp lệ: {value}")
