import openpyxl

from apps.class_sections.application.dto.validate_student_import_dto import (
    ValidateStudentImportDTO,
)


class ClassSectionStudentImportService:

    # =========================================
    # EXCEL STRUCTURE
    # =========================================

    SUBJECT_ROW = 1
    GROUP_ROW = 2
    SEMESTER_ROW = 3
    ACADEMIC_YEAR_ROW = 4

    HEADER_ROW = 5
    STUDENT_START_ROW = 6

    # =========================================
    # READ
    # =========================================

    def read(
        self,
        file,
    ):

        workbook = openpyxl.load_workbook(
            file,
            data_only=True,
        )

        worksheet = workbook.active

        # =========================================
        # VALIDATE HEADER
        # =========================================

        self._validate_header(
            worksheet
        )

        # =========================================
        # READ CLASS SECTION INFO
        # =========================================

        subject_name = self._read_required(
            worksheet=worksheet,
            row=self.SUBJECT_ROW,
            column=2,
            field_name="môn học",
        )

        group_number = self._read_required(
            worksheet=worksheet,
            row=self.GROUP_ROW,
            column=2,
            field_name="nhóm",
        )

        semester_number = self._read_required(
            worksheet=worksheet,
            row=self.SEMESTER_ROW,
            column=2,
            field_name="học kỳ",
        )

        academic_year_name = self._read_required(
            worksheet=worksheet,
            row=self.ACADEMIC_YEAR_ROW,
            column=2,
            field_name="năm học",
        )

        # =========================================
        # PARSE CLASS SECTION INFO
        # =========================================

        subject_name = str(
            subject_name
        ).strip()

        academic_year_name = str(
            academic_year_name
        ).strip()

        group_number = self._parse_integer(
            group_number,
            "Nhóm",
        )

        semester_number = self._parse_integer(
            semester_number,
            "Học kỳ",
        )

        # =========================================
        # VALIDATE VALUES
        # =========================================

        if group_number <= 0:

            raise ValueError(
                "Nhóm phải lớn hơn 0"
            )

        if semester_number not in (
            1,
            2,
            3,
        ):

            raise ValueError(
                "Học kỳ phải là 1, 2 hoặc 3"
            )

        # =========================================
        # READ STUDENTS
        # =========================================

        students = []

        for row in range(
            self.STUDENT_START_ROW,
            worksheet.max_row + 1,
        ):

            student = self._read_student_row(
                worksheet,
                row,
            )

            if student is None:
                continue

            students.append(
                student
            )

        # =========================================
        # CHECK EMPTY
        # =========================================

        if not students:

            raise ValueError(
                "File Excel không có sinh viên"
            )

        return {
            "subject_name": subject_name,

            "group_number": group_number,

            "semester_number": semester_number,

            "academic_year_name": (
                academic_year_name
            ),

            "students": students,
        }

    # =========================================
    # VALIDATE HEADER
    # =========================================

    def _validate_header(
        self,
        worksheet,
    ):

        expected_headers = [
            "MSSV",
            "Họ và tên",
            "Lớp",
        ]

        for index, expected in enumerate(
            expected_headers,
            start=1,
        ):

            actual = worksheet.cell(
                row=self.HEADER_ROW,
                column=index,
            ).value

            if actual is None:

                raise ValueError(
                    f"File Excel thiếu cột '{expected}'"
                )

            actual = str(
                actual
            ).strip()

            if actual != expected:

                raise ValueError(
                    f"Cột {index} phải là "
                    f"'{expected}'"
                )

    # =========================================
    # READ STUDENT ROW
    # =========================================

    def _read_student_row(
        self,
        worksheet,
        row,
    ):

        student_id = worksheet.cell(
            row=row,
            column=1,
        ).value

        full_name = worksheet.cell(
            row=row,
            column=2,
        ).value

        student_class = worksheet.cell(
            row=row,
            column=3,
        ).value

        # -----------------------------------------
        # EMPTY ROW
        # -----------------------------------------

        if (
            student_id is None
            and full_name is None
            and student_class is None
        ):

            return None

        # -----------------------------------------
        # STUDENT ID
        # -----------------------------------------

        if student_id is None:

            raise ValueError(
                f"Dòng {row}: thiếu MSSV"
            )

        student_id = str(
            student_id
        ).strip()

        if not student_id:

            raise ValueError(
                f"Dòng {row}: thiếu MSSV"
            )

        # -----------------------------------------
        # FULL NAME
        # -----------------------------------------

        if full_name is None:

            raise ValueError(
                f"Dòng {row}: "
                f"thiếu họ và tên"
            )

        full_name = str(
            full_name
        ).strip()

        if not full_name:

            raise ValueError(
                f"Dòng {row}: "
                f"thiếu họ và tên"
            )

        # -----------------------------------------
        # STUDENT CLASS
        # -----------------------------------------

        if student_class is None:

            raise ValueError(
                f"Dòng {row}: thiếu lớp"
            )

        student_class = str(
            student_class
        ).strip()

        if not student_class:

            raise ValueError(
                f"Dòng {row}: thiếu lớp"
            )

        # -----------------------------------------
        # DTO
        # -----------------------------------------

        return ValidateStudentImportDTO(
            student_id=student_id,
            full_name=full_name,
            student_class=student_class,
        )

    # =========================================
    # READ REQUIRED
    # =========================================

    def _read_required(
        self,
        worksheet,
        row,
        column,
        field_name,
    ):

        value = worksheet.cell(
            row=row,
            column=column,
        ).value

        if value is None:

            raise ValueError(
                f"File Excel thiếu {field_name}"
            )

        if isinstance(
            value,
            str,
        ):

            value = value.strip()

            if not value:

                raise ValueError(
                    f"File Excel thiếu {field_name}"
                )

        return value

    # =========================================
    # PARSE INTEGER
    # =========================================

    def _parse_integer(
        self,
        value,
        field_name,
    ):

        try:

            return int(value)

        except (
            ValueError,
            TypeError,
        ):

            raise ValueError(
                f"{field_name} không hợp lệ"
            )